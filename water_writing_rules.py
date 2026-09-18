"""Deterministic surface-water monitoring prose and source-basis review."""
import re
from pathlib import Path
from docx import Document

FIXED_LAYOUT_BASIS = "断面垂线和采样点的布设按照国家环保总局颁布的《环境监测技术规范（水和废水部分）》中的规定进行。"


def read_source_context(input_dir):
    evidence, agencies, warnings = [], [], []
    for path in sorted(Path(input_dir).glob("*.docx")):
        try:
            doc = Document(path)
            lines = [p.text for p in doc.paragraphs]
            lines += [" ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        except Exception as exc:
            warnings.append(f"{path.name}：无法读取布设依据，需人工校核（{type(exc).__name__}）。")
            continue
        for line in lines:
            # Only explicitly labelled agencies; never infer them from arbitrary names.
            match = re.search(r"(?:监测单位|检测单位)\s*[：:]\s*([^，。；;\n]+)", line)
            if match and match.group(1).strip() not in agencies:
                agencies.append(match.group(1).strip())
            if re.search(r"断面|垂线|采样点", line) and re.search(r"布设|布置|设置", line):
                for sentence in re.split(r"[。；\n]", line):
                    if not re.search(r"按照|依据|根据|执行|遵循", sentence):
                        continue
                    titles = re.findall(r"《([^》]+)》", sentence)
                    codes = re.findall(r"\bHJ\s*/?T?\s*\d+(?:[—-]\d+)?", sentence)
                    other = [t for t in titles if t.replace("（", "(").replace("）", ")") != "环境监测技术规范(水和废水部分)"]
                    if other or codes:
                        evidence.append(f"{path.name}：{sentence.strip()}")
    return {"agencies": agencies, "conflicts": list(dict.fromkeys(evidence)), "warnings": warnings}


def build_description(table1, factors, table2=None, context=None):
    context = context or {}
    def unique(values):
        return list(dict.fromkeys(str(v).strip() for v in values if v and str(v).strip() not in {"-", "/"}))
    frequencies = unique(row.get("取样频次") for row in table1.get("rows", []))
    dates = unique(row.get("监测时间") for row in (table2 or {}).get("rows", []))
    agencies = unique(context.get("agencies", []))
    subject = "、".join(agencies) if agencies else "本项目"
    timing = "于" + "、".join(dates) if dates else ""
    opening = f"{subject}{timing}开展地表水现状监测"
    if frequencies:
        opening += "，取样频次为" + "；".join(frequencies)
    text = opening + "。" + FIXED_LAYOUT_BASIS
    if factors:
        text += "监测因子包括" + "、".join(factors) + "。"
    return text
