import json
from pathlib import Path
from docx import Document
import surface_water_section_generator as generator
from surface_water_pipeline import parse_standard_config, parse_surface_water_records, evaluate_compliance
from water_writing_rules import FIXED_LAYOUT_BASIS


def test_main_emits_locked_paragraph_and_conflict_review(tmp_path, monkeypatch):
    fixture = json.loads((Path(__file__).parent / 'fixtures/anonymized_eia_case.json').read_text(encoding='utf-8'))
    inputs, outputs = tmp_path / 'input', tmp_path / 'output'
    inputs.mkdir()
    outputs.mkdir()
    for kind in ('plan', 'report'):
        data = fixture[kind]
        doc = Document()
        doc.add_paragraph(data['title'])
        if kind == 'report':
            doc.add_paragraph('监测单位：示例检测有限公司')
        else:
            doc.add_paragraph('断面垂线和采样点的布设依据《项目另行指定规范》执行。')
        table = doc.add_table(rows=1, cols=len(data['headers']))
        for cell, value in zip(table.rows[0].cells, data['headers']):
            cell.text = value
        for values in data['rows']:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        doc.save(inputs / f'{kind}.docx')
    monkeypatch.setenv('EIA_OUTPUT_DIR', str(outputs))
    records = parse_surface_water_records(inputs / 'report.docx')
    standards = parse_standard_config(inputs / 'plan.docx')
    results = evaluate_compliance(records, standards)
    for name, value in [('monitoring_records', records), ('standard_config', standards), ('compliance_results', results)]:
        (outputs / f'{name}.json').write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
    monkeypatch.setattr(generator, 'INPUT_DIR', inputs)
    monkeypatch.setattr(generator, 'OUTPUT_DIR', outputs)
    monkeypatch.setattr(generator, 'DEBUG_DIR', outputs / 'debug_tables')
    monkeypatch.setattr(generator, 'SECTION_PATH', outputs / 'surface_water_section.docx')
    monkeypatch.setattr(generator, 'ENABLE_LLM_TEXT_POLISH', False)
    monkeypatch.setattr(generator, 'build_local_water_status', lambda _: {'text': ''})
    generator.main()
    texts = json.loads((outputs / 'debug_tables/surface_water_section_texts.json').read_text(encoding='utf-8'))
    paragraph = texts['monitoring_time_method_text']
    assert paragraph.startswith('示例检测有限公司于2026-01-01')
    assert paragraph.count(FIXED_LAYOUT_BASIS) == 1
    validation = json.loads((outputs / 'debug_tables/surface_water_formal_text_validation.json').read_text(encoding='utf-8'))
    assert not validation['valid']
    assert any(i['category'] == 'surface_water_layout_basis_review' and '项目另行指定规范' in i['evidence'] for i in validation['issues'])
    word = Document(outputs / 'surface_water_section.docx')
    text = '\n'.join(p.text for p in word.paragraphs)
    assert text.count(FIXED_LAYOUT_BASIS) == 1
    assert paragraph in text
    assert len(word.tables) == 3
    assert '监测时间、频率和方法' in text
