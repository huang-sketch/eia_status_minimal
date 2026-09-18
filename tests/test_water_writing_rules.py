import json
import pytest
from docx import Document
import surface_water_section_generator as generator
from docx_numbering import DocxNumbering
from water_writing_rules import FIXED_LAYOUT_BASIS, build_description, read_source_context


def test_description_order_and_missing_facts():
    text = build_description(
        {"rows": [{"取样频次": "连续三天，每天一次"}]}, ["pH值", "溶解氧"],
        {"rows": [{"监测时间": "2026-09-01"}, {"监测时间": "2026-09-02"}]},
        {"agencies": ["示例检测有限公司"]},
    )
    assert text.startswith("示例检测有限公司于2026-09-01、2026-09-02开展地表水现状监测，取样频次为连续三天，每天一次。")
    assert text.count(FIXED_LAYOUT_BASIS) == 1
    assert text.endswith("监测因子包括pH值、溶解氧。")
    assert text.index("每天一次") < text.index(FIXED_LAYOUT_BASIS) < text.index("监测因子")
    assert build_description({"rows": []}, []) == "本项目开展地表水现状监测。" + FIXED_LAYOUT_BASIS


def test_source_basis_and_agency(tmp_path):
    doc = Document()
    doc.add_paragraph("监测单位：示例检测有限公司")
    doc.add_paragraph("断面垂线和采样点的布设按照《另一布设规范》执行。")
    doc.add_paragraph("pH值检测方法按照《另一实验室规范》执行。")
    doc.add_paragraph(FIXED_LAYOUT_BASIS)
    doc.save(tmp_path / "report.docx")
    result = read_source_context(tmp_path)
    assert result["agencies"] == ["示例检测有限公司"]
    assert len(result["conflicts"]) == 1
    assert "另一布设规范" in result["conflicts"][0]
    assert not result["warnings"]


def test_unreadable_source_produces_review(tmp_path):
    (tmp_path / "bad.docx").write_bytes(b"invalid")
    assert read_source_context(tmp_path)["warnings"]


@pytest.mark.parametrize("mode", ["disabled", "missing_key", "failure", "omitted", "rewritten", "repeated", "normal"])
def test_llm_paths_preserve_paragraph(tmp_path, monkeypatch, mode):
    monkeypatch.setattr(generator, "DEBUG_DIR", tmp_path)
    monkeypatch.setattr(generator, "ENABLE_LLM_TEXT_POLISH", mode != "disabled")
    monkeypatch.setenv("EIA_LLM_API_KEY", "test-key")
    if mode == "missing_key":
        monkeypatch.delenv("EIA_LLM_API_KEY")
    numbering = DocxNumbering(tmp_path / "numbering.json")
    numbering.begin_section("surface_water", "地表水环境现状调查与评价")
    t1 = {"rows": [{"取样频次": "每天一次"}]}
    t2 = {"rows": [{"监测时间": "2026-09-01"}]}
    t3 = {"rows": []}
    for table, key in zip((t1, t2, t3), (generator.SURFACE_TABLE_MONITOR_POINTS, generator.SURFACE_TABLE_MONITOR_RESULTS, generator.SURFACE_TABLE_COMPLIANCE)):
        table.update(table_key=key, caption_suffix=key)
    generator.register_surface_tables(numbering, t1, t2, t3)
    rule = generator.build_rule_texts(t1, t2, ["pH值"], [], "各评价因子达标。", numbering)
    def fake_llm(*args, **kwargs):
        if mode == "failure":
            raise RuntimeError("synthetic unavailable")
        result = dict(rule)
        if mode == "omitted":
            result.pop("monitoring_time_method_text")
        elif mode == "rewritten":
            result["monitoring_time_method_text"] = "按照其他规范布点。监测因子包括氨氮。"
        elif mode == "repeated":
            result["monitoring_time_method_text"] += FIXED_LAYOUT_BASIS
        return result
    monkeypatch.setattr(generator, "chat_completion_json_object_with_recovery", fake_llm)
    result = generator.polish_surface_water_text_with_llm(rule, t1, t2, t3, ["pH值"], [], [], numbering)
    assert result["monitoring_time_method_text"] == rule["monitoring_time_method_text"]
    assert result["monitoring_time_method_text"].count(FIXED_LAYOUT_BASIS) == 1
    assert "检测方法按监测报告" not in result["monitoring_time_method_text"]
    if mode in {"omitted", "rewritten", "repeated", "normal"}:
        audit = json.loads((tmp_path / "surface_water_llm_text_output.json").read_text(encoding="utf-8"))
        assert audit["monitoring_time_method_text"] == rule["monitoring_time_method_text"]
